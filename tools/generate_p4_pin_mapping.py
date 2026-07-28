from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils import get_column_letter


HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEAD_FILL = PatternFill("solid", fgColor="D9EAF7")
OK_FILL = PatternFill("solid", fgColor="E2F0D9")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
RESERVED_FILL = PatternFill("solid", fgColor="F4CCCC")
INFO_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="B7B7B7")


MAPPINGS = [
    # module, signal, H7 pin, P4 GPIO, header, peripheral/direction, status, note, basis
    ("OV2640 DVP", "D0", "PC6", "GPIO2", "P1-3", "DVP input", "已确认", "数据位0；与D1-D7连续布线", "P4 GPIO Matrix + 官方DVP自定义引脚支持"),
    ("OV2640 DVP", "D1", "PC7", "GPIO3", "P1-4", "DVP input", "已确认", "数据位1", "同上"),
    ("OV2640 DVP", "D2", "PC8", "GPIO4", "P1-5", "DVP input", "已确认", "数据位2", "同上"),
    ("OV2640 DVP", "D3", "PC9", "GPIO5", "P1-6", "DVP input", "已确认", "数据位3", "同上"),
    ("OV2640 DVP", "D4", "PC11", "GPIO6", "P1-7", "DVP input", "已确认", "数据位4", "同上"),
    ("OV2640 DVP", "D5", "PD3", "GPIO7", "P1-8", "DVP input", "已确认", "数据位5", "同上"),
    ("OV2640 DVP", "D6", "PB8", "GPIO8", "P1-9", "DVP input", "已确认", "数据位6", "同上"),
    ("OV2640 DVP", "D7", "PB9", "GPIO9", "P1-10", "DVP input", "已确认", "数据位7", "同上"),
    ("OV2640 DVP", "PCLK", "PA6", "GPIO10", "P1-11", "DVP input", "已确认", "时钟线优先短、少过孔；与数据组等长不是硬性要求", "官方DVP控制器支持自定义PCLK"),
    ("OV2640 DVP", "VSYNC", "PB7", "GPIO11", "P1-12", "DVP input", "已确认", "帧同步", "官方DVP控制器支持自定义VSYNC"),
    ("OV2640 DVP", "HREF/DE", "PH8", "GPIO12", "P1-13", "DVP input", "已确认", "ESP-IDF/ESP-Video中对应DE引脚", "官方示例使用DE表示DVP行有效"),
    ("OV2640 DVP", "XCLK(预留)", "当前H7未连接", "GPIO13", "P1-14", "LEDC clock output", "条件确认", "串0欧电阻到摄像头XCLK；现模块若自带时钟则DNP", "现H7工程无XCLK；为摄像头批次兼容预留"),
    ("OV2640 DVP", "SCCB_SCL", "PB4", "GPIO14", "P1-15", "I2C/SCCB open-drain", "已确认", "独立摄像头控制总线，外接上拉按模块情况配置", "官方OV2640驱动使用SCCB(I2C)"),
    ("OV2640 DVP", "SCCB_SDA", "PB3", "GPIO15", "P1-16", "I2C/SCCB open-drain", "已确认", "与语音I2C分开，降低一周迁移风险", "同上"),
    ("OV2640 DVP", "RESET", "PA15", "GPIO16", "P1-17", "GPIO output", "已确认", "低有效，建议10k上拉", "现有OV2640驱动时序"),
    ("OV2640 DVP", "PWDN", "PA0", "GPIO17", "P1-18", "GPIO output", "已确认", "高有效，建议默认安全电平", "现有OV2640驱动时序"),
    ("OV2640 DVP", "FLASH/补光", "PA8", "GPIO18", "P1-19", "GPIO/PWM output", "已确认", "按原模块驱动电路接入，不直接带大电流LED", "H7映射表"),
    ("水箱水位", "TANK_LEVEL", "PB5", "GPIO19", "P1-20", "GPIO input", "已确认", "确认传感器输出不超过3.3V", "当前main.h实际定义"),
    ("按键", "冷水按钮", "PF7", "GPIO20", "P1-21", "GPIO input", "已确认", "低有效，启用上拉并做消抖", "当前camera_app.c实际使用"),
    ("按键", "热水按钮", "PF6", "GPIO21", "P1-22", "GPIO input", "条件确认", "核心板同时连MIPI-DSI触摸INT；不用DSI屏时可用", "P4小系统板IO表"),
    ("饮水机照明", "AMBIENT_LED1", "PH5", "GPIO22", "P1-23", "GPIO/PWM output", "已确认", "经原驱动级，不直接带大电流负载", "当前main.h实际定义"),
    ("语音模块", "Voice_SCL", "PD12", "GPIO32", "P1-24", "I2C SCL", "条件确认", "板载MIPI侧电平转换及上拉；不用MIPI摄像头时可用，实测总上拉", "核心板原理图 IIC_SCL"),
    ("语音模块", "Voice_SDA", "PD13", "GPIO33", "P1-25", "I2C SDA", "条件确认", "与GPIO32同一总线约束", "核心板原理图 IIC_SDA"),
    ("OLED", "SCK/SCL", "PB13", "GPIO30", "P2-12", "SPI clock output", "已确认", "沿用SPI命名；建议降低首版时钟后再提速", "GPIO30独立"),
    ("OLED", "MOSI/SDA", "PB15", "GPIO29", "P2-8", "SPI MOSI output", "已确认", "仅输出", "GPIO29独立"),
    ("OLED", "CS", "PC1", "GPIO28", "P2-9", "GPIO/SPI CS output", "已确认", "默认拉高避免上电误选中", "GPIO28独立"),
    ("OLED", "RES", "PC4", "GPIO27", "P2-10", "GPIO output", "已确认", "保持原复位时序", "GPIO27独立"),
    ("OLED", "DC", "PC5", "GPIO26", "P2-11", "GPIO output", "已确认", "数据/命令选择", "GPIO26独立"),
    ("温度传感器", "DS18B20_DQ", "PE3", "GPIO31", "P2-7", "1-Wire open-drain", "已确认", "外接4.7k上拉到3.3V", "GPIO31独立"),
    ("泵控", "PUMP_PWM", "PH6", "GPIO36", "P2-14", "LEDC/MCPWM output", "已确认", "上电默认关泵；必须经原MOSFET/驱动级", "GPIO36独立"),
    ("AS608", "模块RX / 主控TX", "PA2", "GPIO46", "P2-16", "UART TX", "已确认", "P4发送到AS608接收", "GPIO46独立"),
    ("AS608", "模块TX / 主控RX", "PA3", "GPIO47", "P2-17", "UART RX", "已确认", "AS608发送到P4接收", "GPIO47独立"),
    ("AS608", "WAK", "PE6", "GPIO48", "P2-18", "GPIO input", "已确认", "按模块有效电平配置上下拉", "GPIO48独立"),
    ("外置C6", "P4_TX -> C6_RX", "PB10/U3_TX", "GPIO49", "P2-20", "UART TX", "已确认", "保持现有串口协议和波特率；对端仍接C6 GPIO5/RX", "GPIO49独立"),
    ("外置C6", "P4_RX <- C6_TX", "PB11/U3_RX", "GPIO50", "P2-19", "UART RX", "已确认", "保持现有串口协议和波特率；对端仍接C6 GPIO4/TX", "GPIO50独立"),
    ("饮水机照明", "AMBIENT_LED2", "PH2", "GPIO23", "P2-22", "GPIO/PWM output", "已确认", "经原驱动级", "当前main.h实际定义"),
    ("调试串口", "UART0_TX", "PA9/USART1_TX", "GPIO37", "P1-1", "UART0 TX", "已确认", "核心板跳线连接CH343；调试日志专用", "核心板原理图P3"),
    ("调试串口", "UART0_RX", "PA10/USART1_RX", "GPIO38", "P1-2", "UART0 RX", "已确认", "核心板跳线连接CH343", "核心板原理图P3"),
    ("状态指示", "LED0/启动状态", "PB1", "GPIO51", "P2-21", "GPIO output", "条件确认", "核心板已接红色LED，优先仅作板载状态灯", "核心板原理图LED"),
    ("状态指示", "LED1", "PB0", "GPIO53", "P2-3", "GPIO output", "条件确认", "与LCD_BL相连；不用MIPI/DSI屏时可使用", "核心板IO表标注可独立"),
]


P1 = {
    1: "GPIO37", 2: "GPIO38", 3: "GPIO2", 4: "GPIO3", 5: "GPIO4", 6: "GPIO5",
    7: "GPIO6", 8: "GPIO7", 9: "GPIO8", 10: "GPIO9", 11: "GPIO10", 12: "GPIO11",
    13: "GPIO12", 14: "GPIO13", 15: "GPIO14", 16: "GPIO15", 17: "GPIO16", 18: "GPIO17",
    19: "GPIO18", 20: "GPIO19", 21: "GPIO20", 22: "GPIO21", 23: "GPIO22", 24: "GPIO32",
    25: "GPIO33",
}

P2 = {
    1: "5V", 2: "GND", 3: "GPIO53", 4: "GPIO52", 5: "GPIO24", 6: "GPIO25",
    7: "GPIO31", 8: "GPIO29", 9: "GPIO28", 10: "GPIO27", 11: "GPIO26", 12: "GPIO30",
    13: "GPIO35", 14: "GPIO36", 15: "GPIO45", 16: "GPIO46", 17: "GPIO47", 18: "GPIO48",
    19: "GPIO50", 20: "GPIO49", 21: "GPIO51", 22: "GPIO23", 23: "RESET", 24: "USB_DM",
    25: "USB_DP",
}


def style_sheet(ws, widths: list[int]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 34
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if cell.row > 1:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.45, bottom=0.45, header=0.2, footer=0.2)
    ws.print_title_rows = "1:1"
    ws.print_area = ws.dimensions


def add_mapping_sheet(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "H7到P4映射"
    ws.append(["序号", "模块", "信号", "STM32H7引脚", "ESP32-P4 GPIO", "核心板针位", "P4外设/方向", "状态", "电气与布板说明", "依据"])
    for idx, row in enumerate(MAPPINGS, start=1):
        ws.append([idx, *row])
        status_cell = ws.cell(ws.max_row, 8)
        status_cell.fill = OK_FILL if row[6] == "已确认" else WARN_FILL
    style_sheet(ws, [7, 15, 22, 18, 16, 14, 22, 12, 48, 42])
    ws.sheet_view.showGridLines = False


def add_pin_audit_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("P4排针资源审计")
    ws.append(["核心板针位", "网络/GPIO", "本方案分配", "资源状态", "板载连接/注意事项"])
    assigned = {row[4]: f"{row[0]} - {row[1]}" for row in MAPPINGS}
    notes = {
        "GPIO21": "同时连接MIPI-DSI触摸INT；不用DSI屏时可用",
        "GPIO32": "板载I2C并接MIPI侧1.8V电平转换",
        "GPIO33": "板载I2C并接MIPI侧1.8V电平转换",
        "GPIO24": "USB Serial/JTAG D-，保留",
        "GPIO25": "USB Serial/JTAG D+，保留",
        "GPIO35": "BOOT，必须保留启动功能",
        "GPIO37": "UART0 TX，可通过P3跳线连CH343",
        "GPIO38": "UART0 RX，可通过P3跳线连CH343",
        "GPIO45": "MIPI-DSI触摸复位；本方案暂不使用",
        "GPIO51": "核心板红色LED",
        "GPIO52": "MIPI-DSI LCD复位；本方案保留",
        "GPIO53": "LCD背光；不用DSI屏时可用",
        "RESET": "芯片硬件复位，禁止改作信号",
        "USB_DM": "USB2.0 HS专用信号，保留",
        "USB_DP": "USB2.0 HS专用信号，保留",
        "5V": "核心板电源输入；泵电源不要从此针直接取",
        "GND": "数字地；泵功率回流应单独规划后单点汇接",
    }
    hard_reserved = {"GPIO24", "GPIO25", "GPIO35", "RESET", "USB_DM", "USB_DP"}
    for header, pins in (("P1", P1), ("P2", P2)):
        for pin_num, net in pins.items():
            location = f"{header}-{pin_num}"
            allocation = assigned.get(location, "")
            if allocation:
                state = "已分配"
                fill = OK_FILL
            elif net in hard_reserved:
                state = "保留/禁用"
                fill = RESERVED_FILL
            elif net in {"5V", "GND"}:
                state = "电源"
                fill = INFO_FILL
            else:
                state = "备用"
                fill = WARN_FILL
            ws.append([location, net, allocation, state, notes.get(net, "")])
            ws.cell(ws.max_row, 4).fill = fill
    style_sheet(ws, [15, 16, 34, 15, 52])
    ws.sheet_view.showGridLines = False


def add_dvp_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("OV2640_DVP说明")
    ws.append(["项目", "结论/参数", "设计动作"])
    rows = [
        ("总体结论", "ESP32-P4可继续使用现有OV2640 DVP，不需要因引脚不足改MIPI", "先按DVP打样并保留MIPI作为后备"),
        ("官方驱动", "esp-video-components支持ESP32-P4 DVP和OV2640；含320x240 JPEG配置", "ESP-IDF v5.5.2，使用Custom development board配置"),
        ("数据总线", "D0-D7 = GPIO2-GPIO9", "放同一侧、连续走线、少过孔，串阻位靠近P4预留22-33欧"),
        ("同步信号", "PCLK=GPIO10，VSYNC=GPIO11，HREF/DE=GPIO12", "PCLK优先最短；确认驱动极性与现有OV2640寄存器配置"),
        ("传感器时钟", "当前H7未输出XCLK，推测模块自带时钟；GPIO13作为可选XCLK", "加0欧/DNP选择位，并在摄像头座保留XCLK脚"),
        ("控制总线", "SCCB SCL/SDA = GPIO14/GPIO15", "3.3V开漏；根据模块现有上拉决定底板是否装上拉"),
        ("控制脚", "RESET=GPIO16，PWDN=GPIO17，补光=GPIO18", "上电时先保持摄像头安全态，随后按现固件时序释放"),
        ("帧格式", "现工程为320x240 JPEG并软件解码", "P4首版继续JPEG以减少模型输入变化；后续可评估YUV/RGB直采"),
        ("供电", "摄像头与P4均为3.3V逻辑，仍需核对具体OV2640模块电源脚", "摄像头电源旁放10uF+0.1uF；补光/泵负载不得污染3.3V"),
        ("禁止占用", "USB、BOOT、Flash、TF和MIPI专用引脚不用于DVP", "按P4排针资源审计表布线"),
        ("MIPI后备", "核心板已自带MIPI-CSI接口，但更换传感器会引入模型域偏移", "只有DVP实机采集失败且驱动无法修复时再切MIPI并重采数据"),
    ]
    for row in rows:
        ws.append(row)
    style_sheet(ws, [22, 62, 62])
    ws.sheet_view.showGridLines = False


def add_power_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("电源与对端连接")
    ws.append(["对象", "原连接", "P4底板连接", "是否占GPIO", "设计说明"])
    rows = [
        ("核心板供电", "STM32底板主电源", "5V -> P2-1；GND -> P2-2及其他GND", "否", "核心板自带5V转3.3V DCDC；不要把泵电流从P2-1穿过核心板"),
        ("AS608 Vi", "3.3V", "底板稳压3.3V", "否", "确认模块Vi/Vt定义和峰值电流，旁路0.1uF+10uF"),
        ("AS608 Vt", "3.3V", "底板稳压3.3V", "否", "沿用原模块接法"),
        ("AS608 GND", "GND", "数字GND", "否", "与P4共地"),
        ("外置C6供电", "现有C6模块电源", "继续使用现有3.3V供电支路", "否", "P4和C6必须共地；不要迁移C6联网电路"),
        ("外置C6 RX", "C6 GPIO5/RX", "P4 GPIO49/P2-20 TX", "GPIO49", "交叉连接，保持现有电平和波特率"),
        ("外置C6 TX", "C6 GPIO4/TX", "P4 GPIO50/P2-19 RX", "GPIO50", "交叉连接"),
        ("冷水按钮蓝色LED", "5V/GND", "保留原5V限流/灯电路", "否", "不是P4信号，不能直接接GPIO"),
        ("热水按钮红色LED", "5V/GND", "保留原5V限流/灯电路", "否", "不是P4信号，不能直接接GPIO"),
        ("泵功率电源", "原泵电源", "继续走独立功率支路", "GPIO36仅控制", "泵回流与数字地单点汇接，保留续流/TVS/大电容"),
        ("OV2640电源", "原摄像头供电", "按模块额定电压保留", "否", "数据电平必须为3.3V；电源脚不能仅凭模块名称判断"),
    ]
    for row in rows:
        ws.append(row)
    style_sheet(ws, [24, 27, 42, 18, 62])
    ws.sheet_view.showGridLines = False


def add_source_sheet(wb: Workbook, source_h7: Path, source_p4: Path) -> None:
    ws = wb.create_sheet("依据与假设")
    ws.append(["类型", "内容"])
    rows = [
        ("H7输入表", str(source_h7)),
        ("P4官方IO表", str(source_p4)),
        ("P4原理图", "DNESP32P4M V1.0 SCH.pdf，重点核对P1/P2、POWER、DEVICE页"),
        ("P4硬件手册", "DNESP32P4硬件参考手册_V1.0.pdf，第1.3-1.5及2.2节"),
        ("乐鑫驱动依据", "https://github.com/espressif/esp-video-components"),
        ("映射原则", "保留核心板Flash/PSRAM/USB/BOOT/TF/MIPI专用信号；外置C6继续UART联网"),
        ("重要假设", "最终使用自制底板而不是正点原子原配底板，因此RGB LCD、I2S等原配底板占用不适用"),
        ("封装方向", "必须使用资料包官方PcbLib/IntLib建立P1/P2封装，不可按照片手工镜像排针"),
        ("待实机确认", "OV2640模块是否自带XCLK、DVP极性、I2C总上拉、各外设电平、泵上电默认关断"),
    ]
    for row in rows:
        ws.append(row)
    style_sheet(ws, [22, 110])
    ws.sheet_view.showGridLines = False


def copy_source_sheet(wb: Workbook, source_h7: Path) -> None:
    src_wb = load_workbook(source_h7, data_only=False)
    src_ws = src_wb[src_wb.sheetnames[0]]
    ws = wb.create_sheet("原始H7表")
    for row in src_ws.iter_rows():
        ws.append([cell.value for cell in row])
    for idx in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 28
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.sheet_view.showGridLines = False


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--h7", required=True, type=Path)
    parser.add_argument("--p4", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    wb = Workbook()
    add_mapping_sheet(wb)
    add_pin_audit_sheet(wb)
    add_dvp_sheet(wb)
    add_power_sheet(wb)
    add_source_sheet(wb, args.h7, args.p4)
    copy_source_sheet(wb, args.h7)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
