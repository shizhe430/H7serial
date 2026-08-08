import openpyxl
wb = openpyxl.load_workbook(r'C:\Users\17360\Desktop\STM32H7到ESP32P4引脚映射_已核对.xlsx')
ws = wb.active
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
    print(row)
